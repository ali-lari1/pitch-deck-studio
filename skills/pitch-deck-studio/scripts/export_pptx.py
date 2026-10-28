#!/usr/bin/env python3
"""Package a built deck as deck.pptx: one full-bleed slide image per slide, with speaker notes.

Usage:
    python3 export_pptx.py <project_dir>

The result opens in PowerPoint, Keynote and Google Slides with notes intact, but the slide text is NOT editable.
For an editable deck, hand deck.md + brand.json to a PPTX skill (see references/compose.md).
Needs `python-pptx` (pip install python-pptx) and slide PNGs from `export.py <project> --png`
(they are rendered automatically if missing).
"""
import html
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))


def notes_from_html(path):
    with open(path, encoding="utf-8") as f:
        doc = f.read()
    out = []
    for m in re.finditer(r'<section class="slide[^"]*"[^>]*>.*?</section>(<aside class="notes">(.*?)</aside>)?', doc, re.S):
        raw = m.group(2) or ""
        txt = re.sub(r"</p>\s*<p>", "\n\n", raw)
        txt = html.unescape(re.sub(r"<[^>]+>", "", txt)).strip()
        out.append(txt)
    return out


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help"):
        print(__doc__)
        return 0
    root = os.path.abspath(argv[1])
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print("python-pptx is not installed. Run: pip install python-pptx  (or use a PPTX skill; see references/compose.md)",
              file=sys.stderr)
        return 2
    html_path = os.path.join(root, "deck.html")
    if not os.path.exists(html_path):
        print("no deck.html - run build.py first", file=sys.stderr)
        return 1
    notes = notes_from_html(html_path)
    sdir = os.path.join(root, "slides")
    pngs = sorted(f for f in os.listdir(sdir)) if os.path.isdir(sdir) else []
    if len(pngs) != len(notes):
        subprocess.run([sys.executable, os.path.join(HERE, "export.py"), root, "--png"], check=False)
        pngs = sorted(f for f in os.listdir(sdir) if f.endswith(".png"))
    if len(pngs) != len(notes):
        print("expected %d slide PNGs, found %d" % (len(notes), len(pngs)), file=sys.stderr)
        return 3
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]
    for png, note in zip(pngs, notes):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(sdir, png), 0, 0, prs.slide_width, prs.slide_height)
        if note:
            s.notes_slide.notes_text_frame.text = note
    out = os.path.join(root, "deck.pptx")
    prs.save(out)
    print("pptx:  %s (%d slides, image-based, notes included)" % (out, len(pngs)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
